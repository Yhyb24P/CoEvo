"""Independent bookkeeping checks; does not infer Agent learning or scientific correctness."""
import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def check(original, submitted, results):
    if not Path(submitted).is_file():
        return {'file_present': False, 'mechanical_checks_passed': False}
    before, after = load_workbook(original), load_workbook(submitted)
    problems = []
    if before.sheetnames != after.sheetnames:
        problems.append('Sheet names/order changed')
    for old in before:
        if old.title not in after:
            continue
        new = after[old.title]
        for row in old:
            for cell in row:
                target = new[cell.coordinate]
                edited = old.title == 'Runs' and cell.row in (3, 4) and cell.column in range(2, 7)
                if not edited and target.value != cell.value:
                    problems.append(f'Unrequested value/formula change: {old.title}!{cell.coordinate}')
                style = lambda c: (str(c.font), str(c.fill), str(c.border), str(c.alignment), c.number_format, str(c.protection))
                if style(target) != style(cell):
                    problems.append(f'Style changed: {old.title}!{cell.coordinate}')
                comment = lambda c: (c.comment.text, c.comment.author) if c.comment else None
                if comment(cell) != comment(target):
                    problems.append(f'Comment changed: {old.title}!{cell.coordinate}')
        if str(old.merged_cells) != str(new.merged_cells) or old.freeze_panes != new.freeze_panes or old.auto_filter.ref != new.auto_filter.ref:
            problems.append(f'Structure changed: {old.title}')
        if {k: d.width for k, d in old.column_dimensions.items()} != {k: d.width for k, d in new.column_dimensions.items()}:
            problems.append(f'Column widths changed: {old.title}')
    actual_rows = {}
    if 'Runs' in after:
        sheet = after['Runs']
        for case, row in [('software', 3), ('science', 4)]:
            values = [sheet.cell(row, col).value for col in range(1, 7)]
            actual_rows[case] = values
            expected = results['cases'][case]
            if not isinstance(values[1], (float, int)) or not math.isclose(values[1], expected['wall_seconds'], abs_tol=.1):
                problems.append(f'Incorrect/missing seconds: {case}')
            if values[2] != expected['reported_tool_counter']:
                problems.append(f'Incorrect/missing tool counter: {case}')
            if not all(str(values[col] or '').strip() for col in (3, 4, 5)):
                problems.append(f'Stop/outcome/notes missing: {case}')
    return {'file_present': True, 'mechanical_checks_passed': not problems,
            'problems': problems, 'submitted_rows': actual_rows,
            'new_task_minutes_reference': sum(x['wall_seconds'] for x in results['cases'].values()) / 60,
            'review_pending': ['Outcome and stop labels correctly distinguish execution/reproduction from task/scientific success',
                               'Software counter note explains the blocked budget attempt',
                               'Delivery explanation and workbook calculations match source records'],
            'scope': 'Workbook requirement evidence; semantic outcome review remains necessary'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--original', required=True)
    parser.add_argument('--submitted', required=True)
    parser.add_argument('--results', required=True)
    args = parser.parse_args()
    print(json.dumps(check(args.original, args.submitted, json.loads(Path(args.results).read_text())), ensure_ascii=False, indent=2))
