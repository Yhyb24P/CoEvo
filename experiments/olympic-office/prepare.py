"""Freeze a complete spreadsheet task and native-history comparison using existing Pi execution."""
import argparse
import hashlib
import json
import runpy
import shutil
from pathlib import Path

from coevo.pi_office import prepare as prepare_office
from coevo.project import hashes
from coevo.runtime import write_json
from coevo.sandbox import command

ROOT=Path(__file__).resolve().parents[2]


def make_template(data, destination):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    reference=runpy.run_path(str(ROOT/'experiments/olympic-office/check.py'))['reference']
    _,summary,_=reference(data)
    w=Workbook();s=w.active;s.title='Summary'
    s.append(['Olympic dataset review']);s.merge_cells('A1:H1')
    s.append(['Year','NOCs','Gold','Silver','Bronze','Total','Events (programs)','Host'])
    s.append(summary[2016]);s.append([None]*8);s.append([None]*8)
    s['A7']='Reviewed medal totals';s['F7']='=SUM(F3:F5)'
    s['F7'].comment=Comment('Keep the formula. Dataset medal totals are not counts of sporting events.','CoEvo')
    s.freeze_panes='A3';s.auto_filter.ref='A2:H5'
    for row in s.iter_rows(min_row=2,max_row=5,max_col=8):
        for c in row:
            c.font=Font(name='Calibri',size=11,bold=c.row==2)
            if c.row==2:c.fill=PatternFill('solid',fgColor='D9EAF7')
    for col in range(1,9):s.column_dimensions[get_column_letter(col)].width=18 if col<8 else 85
    a=w.create_sheet('Archive');a.append(['Frozen historical record','2016 source snapshot'])
    a.append(['2016 NOCs',summary[2016][1]]);a['B2'].comment=Comment('Historical input; do not overwrite.','CoEvo')
    w.save(destination)


def prepare(out, data, parent, source, pi_entry, environment, skill_cache):
    out,data,parent=map(lambda p:Path(p).resolve(),(out,data,parent))
    out.mkdir(parents=True,exist_ok=False)
    inputs=out/'inputs';inputs.mkdir()
    make_template(data, inputs/'olympic-audit.xlsx')
    prompt=(ROOT/'experiments/olympic-office/task.txt').read_text()
    parent_trace=json.loads((parent/'transcript.json').read_text())
    parent_session=Path(parent_trace['session_file'])
    for arm in ['R0','R1-native']:
        root=out/arm
        request=prepare_office(root,source,pi_entry,environment,skill_cache)
        workspace=root/'workspace'
        # Replace only this fresh preparation's example inputs, before execution.
        for name in ['research-ledger.xlsx','first-W0-results.json']:(workspace/name).unlink()
        shutil.copy2(inputs/'olympic-audit.xlsx',workspace/'olympic-audit.xlsx')
        (workspace/'data').mkdir()
        for name in ['summerOly_medal_counts.csv','summerOly_hosts.csv','summerOly_programs.csv']:
            shutil.copy2(data/name,workspace/'data'/name)
        readonly=['.skills','data','olympic-audit.xlsx']
        prefix=command(workspace,environment,[],readonly_paths=readonly)[:-1]
        prefix+=['--ro-bind',str(Path(skill_cache).resolve()),'/opt/skill-cache','--setenv','SPREADSHEET_SKILL_CACHE','/opt/skill-cache','--']
        request.update(prompts=[prompt],readonly_paths=readonly,sandbox_prefix=prefix,
                       timeout_ms=1500000,max_tool_calls=80,repeat_failure_limit=4)
        if arm=='R1-native':request['resume_file']=str(parent_session)
        write_json(root/'agent_request.json',request)
        task=json.loads((root/'task.json').read_text())
        task.update(source_cluster='mcm-2025-c',split='D',train_authorized=False,
                    case='olympic-office',arm=arm,feedback_assisted=False,
                    experience_origin='native prior Office recovery' if arm=='R1-native' else None,
                    prior_experience_feedback_assisted=(arm=='R1-native'),
                    inputs={p:h for p,h in hashes(workspace).items() if not p.startswith('.skills/')},
                    initial_files=hashes(workspace),
                    budget={'wall_seconds':1500,'max_tool_calls':80,'max_output_tokens':32768})
        write_json(root/'task.json',task)
    initial=json.loads((out/'R0/task.json').read_text())['initial_files']
    assert initial==json.loads((out/'R1-native/task.json').read_text())['initial_files']
    sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    manifest={'status':'prepared','source_cluster':'mcm-2025-c','split':'D','arms':['R0','R1-native'],
              'parent_run':str(parent),'parent_trace_sha256':sha(parent/'transcript.json'),
              'parent_session':str(parent_session),'parent_session_sha256':sha(parent_session),
              'parent_message_count':len(parent_trace['messages']),
              'data_sha256':{p.name:sha(p) for p in (out/'R0/workspace/data').iterdir()},
              'prompt_sha256':sha(ROOT/'experiments/olympic-office/task.txt'),
              'checker_sha256':sha(ROOT/'experiments/olympic-office/check.py'),
              'same_initial_files':True,'same_model_weights':'Qwen3.5-4B W0','order':['R0','R1-native'],
              'scope':'Single complete D task; raw native prior experience versus fresh session. Context and inference costs intentionally differ. No H or isolated memory-summary claim.'}
    write_json(out/'comparison.json',manifest)
    return manifest


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',required=True);p.add_argument('--data',required=True);p.add_argument('--parent',required=True)
    p.add_argument('--source',default='.cache/pilotdeck');p.add_argument('--pi-entry',default='.cache/pi-1.2/node_modules/@earendil-works/pi-coding-agent/dist/index.js')
    p.add_argument('--environment',default='.cache/pilotdeck-office-env');p.add_argument('--skill-cache',default='.cache/spreadsheets-skill-cache')
    a=p.parse_args();print(json.dumps(prepare(a.out,a.data,a.parent,a.source,a.pi_entry,a.environment,a.skill_cache),ensure_ascii=False))
