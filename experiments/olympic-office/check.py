"""Independent file/data checks for the complete Olympic Office task."""
import csv
import json
from collections import Counter
from pathlib import Path
from zipfile import BadZipFile


def clean(s):
    return ' '.join(str(s).replace('\ufeff','').split())


def reference(data):
    # Reviewed source encodings; the Agent chooses its own read strategy.
    def read(name, encoding):
        with (Path(data)/name).open(encoding=encoding,newline='') as f:
            return list(csv.DictReader(f))
    raw=read('summerOly_medal_counts.csv','utf-8-sig')
    medals=[{k:(clean(r[k]) if k=='NOC' else int(r[k])) for k in ['Year','NOC','Rank','Gold','Silver','Bronze','Total']} for r in raw]
    hosts={int(r['Year']):clean(r['Host']) for r in read('summerOly_hosts.csv','utf-8-sig')}
    programs=read('summerOly_programs.csv','cp1252')
    total,= [r for r in programs if r['Sport']=='Total events' and r['Discipline']=='Total events']
    summary={}
    for year in (2016,2020,2024):
        rows=[r for r in medals if r['Year']==year]
        summary[year]=[year,len({r['NOC'] for r in rows}),*[sum(r[k] for r in rows) for k in ['Gold','Silver','Bronze','Total']],int(total[str(year)]),hosts[year]]
    by_year={y:{r['NOC']:r['Total'] for r in medals if r['Year']==y} for y in (2020,2024)}
    comparison={n:(by_year[2020].get(n,0),by_year[2024].get(n,0)) for n in by_year[2020].keys()|by_year[2024].keys()}
    return medals,summary,comparison


def check(workspace):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from openpyxl.utils import get_column_letter
    root=Path(workspace); problems=[]
    expected,summary,comparison=reference(root/'data')
    csv_path=root/'clean_medals.csv'; output=root/'updated-olympic-audit.xlsx'
    if not csv_path.is_file():
        problems.append('clean_medals.csv missing')
    else:
        try:
            with csv_path.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
            actual=[{k:(r[k] if k=='NOC' else int(r[k])) for k in expected[0]} for r in rows]
            key=lambda r:tuple(r[k] for k in expected[0])
            if Counter(map(key,actual))!=Counter(map(key,expected)):
                problems.append('Clean medal records differ from source facts or normalization')
        except (KeyError,ValueError,UnicodeError):
            problems.append('Clean medal schema/encoding/numbers invalid')
    if not output.is_file():
        problems.append('updated-olympic-audit.xlsx missing')
    else:
        old=load_workbook(root/'olympic-audit.xlsx')
        try:
            new=load_workbook(output)
        except (BadZipFile,InvalidFileException,IndexError,ValueError,KeyError) as error:
            problems.append('Submitted workbook cannot be read: '+type(error).__name__+': '+str(error))
            problems += [name+' missing/empty' for name in ('audit.md','prepare_audit.py')
                         if not (root/name).is_file() or not (root/name).stat().st_size]
            return {'file_checks_passed':False,'problems':problems,'source_rows':len(expected),
                    'workbook_readable':False,'expected_summary':summary,'comparison_nocs':len(comparison),
                    'scope':'Unreadable submitted workbook; raw file retained. Other scientific claims remain unverified.'}
        if new.sheetnames[:len(old.sheetnames)]!=old.sheetnames:
            problems.append('Original sheet order changed')
        for sheet in old:
            if sheet.title not in new:
                problems.append('Original sheet missing: '+sheet.title);continue
            dst=new[sheet.title]
            for row in sheet:
                for c in row:
                    d=dst[c.coordinate]
                    allowed=sheet.title=='Summary' and c.row in (4,5) and c.column<=8
                    if not allowed and c.value!=d.value:problems.append('Protected value/formula changed: '+sheet.title+'!'+c.coordinate)
                    style=lambda c:tuple(str(getattr(c,k)) for k in ('font','fill','border','alignment','number_format','protection'))
                    if style(c)!=style(d):problems.append('Style changed: '+sheet.title+'!'+c.coordinate)
                    note=lambda x:(x.comment.text,x.comment.author) if x.comment else None
                    if note(c)!=note(d):problems.append('Comment changed: '+sheet.title+'!'+c.coordinate)
            if str(sheet.merged_cells)!=str(dst.merged_cells) or sheet.freeze_panes!=dst.freeze_panes or sheet.auto_filter.ref!=dst.auto_filter.ref:
                problems.append('Original structure changed: '+sheet.title)
            if {k:v.width for k,v in sheet.column_dimensions.items()}!={k:v.width for k,v in dst.column_dimensions.items()}:
                problems.append('Column widths changed: '+sheet.title)
        if 'Summary' in new:
            for row,year in [(4,2020),(5,2024)]:
                got=[new['Summary'].cell(row,col).value for col in range(1,9)]
                if got!=summary[year]:problems.append('Summary differs: '+str(year))
        if 'Comparison' not in new:problems.append('Comparison missing')
        else:
            rows=list(new['Comparison'].values)
            try:
                headers=[str(x) for x in rows[0]]
                indices=[headers.index(k) for k in ['NOC','Total2020','Total2024','Change']]
                actual={}
                for excel_row,row in enumerate(rows[1:],2):
                    if row[indices[0]] is None:continue
                    n,a,b,d=[row[i] for i in indices]
                    if isinstance(d,str) and d.startswith('='):
                        formula=f'={get_column_letter(indices[2]+1)}{excel_row}-{get_column_letter(indices[1]+1)}{excel_row}'
                        if d.replace('$','').replace(' ','').upper()==formula:
                            d=b-a if all(isinstance(v,(int,float)) for v in (a,b)) else None
                    if n in actual or not all(isinstance(v,(int,float)) for v in (a,b,d)) or d!=b-a:
                        raise ValueError('Duplicate/non-numeric/inconsistent comparison')
                    actual[n]=(a,b)
                if actual!=comparison:problems.append('Comparison differs from union of source records')
            except (IndexError,ValueError,TypeError):problems.append('Comparison schema/data invalid')
        if 'Quality' not in new or new['Quality'].max_row<2:problems.append('Quality sheet missing/empty')
    for name in ('audit.md','prepare_audit.py'):
        if not (root/name).is_file() or not (root/name).stat().st_size:problems.append(name+' missing/empty')
    return {'file_checks_passed':not problems,'problems':problems,'source_rows':len(expected),'expected_summary':summary,
            'comparison_nocs':len(comparison),'semantic_review_pending':['Quality findings and explanations','Absence convention and source attribution','Reproduction from original inputs','Truthful final delivery'],
            'scope':'File/data evidence; semantic review and original-input preservation are separate'}


if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--workspace',required=True);a=p.parse_args()
    print(json.dumps(check(a.workspace),ensure_ascii=False,indent=2))
